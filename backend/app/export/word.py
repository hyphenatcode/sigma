"""Word (.docx) export via python-docx (§3.8).

Walks the same `ReportContent` / `APATable` objects the HTML path uses, so the
Word file carries the same numbers, the same table structure and the same note
rows as the preview and the PDF.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.apa.tables import APATable
from app.export.html import ReportContent

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
TABLE_SIZE = Pt(10.5)
NOTE_SIZE = Pt(9.5)


def _set_cell_border(cell, *, top: bool = False, bottom: bool = False) -> None:
    """APA tables use horizontal rules only — no vertical or grid lines."""
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, enabled in (("top", top), ("bottom", bottom)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if enabled else "nil")
        if enabled:
            element.set(qn("w:sz"), "6")
            element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)


def _add_apa_table(document: Document, table: APATable) -> None:
    number = document.add_paragraph()
    number_run = number.add_run(table.label)
    number_run.bold = True
    number_run.font.name = BODY_FONT
    number_run.font.size = BODY_SIZE

    title = document.add_paragraph()
    title_run = title.add_run(table.title)
    title_run.italic = True
    title_run.font.name = BODY_FONT
    title_run.font.size = BODY_SIZE

    word_table = document.add_table(rows=1, cols=len(table.columns))
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    alignment = table.align or (["left"] + ["right"] * (len(table.columns) - 1))

    header_cells = word_table.rows[0].cells
    for index, (heading, align) in enumerate(zip(table.columns, alignment)):
        cell = header_cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.LEFT
        )
        run = paragraph.add_run(heading)
        run.font.name = BODY_FONT
        run.font.size = TABLE_SIZE
        _set_cell_border(cell, top=True, bottom=True)

    for row_index, row in enumerate(table.rows):
        cells = word_table.add_row().cells
        is_last = row_index == len(table.rows) - 1
        for index, (value, align) in enumerate(zip(row, alignment)):
            cell = cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(str(value))
            run.font.name = BODY_FONT
            run.font.size = TABLE_SIZE
            _set_cell_border(cell, bottom=is_last)

    if table.note:
        note = document.add_paragraph()
        prefix = note.add_run("Not. ")
        prefix.italic = True
        prefix.font.name = BODY_FONT
        prefix.font.size = NOTE_SIZE
        body = note.add_run(table.note)
        body.font.name = BODY_FONT
        body.font.size = NOTE_SIZE

    document.add_paragraph()


def render_docx(content: ReportContent) -> bytes:
    """Build the .docx and return its bytes."""
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run(content.title)
    heading_run.bold = True
    heading_run.font.name = BODY_FONT
    heading_run.font.size = Pt(14)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        f"{content.dataset_name} · {content.generated_at.strftime('%d.%m.%Y %H:%M')}"
    )
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    section = document.add_paragraph()
    section_run = section.add_run("Bulgular")
    section_run.bold = True
    section_run.font.name = BODY_FONT
    section_run.font.size = BODY_SIZE

    body = document.add_paragraph(content.interpretation)
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if content.substitution_note:
        note = document.add_paragraph()
        label = note.add_run("Analiz seçimi: ")
        label.bold = True
        label.font.size = Pt(10.5)
        text = note.add_run(content.substitution_note)
        text.font.size = Pt(10.5)

    document.add_paragraph()
    for table in content.tables:
        _add_apa_table(document, table)

    footer = document.add_paragraph()
    footer_run = footer.add_run(
        "Bu rapor Sigma tarafından oluşturulmuştur. Tüm istatistiksel değerler "
        "deterministik olarak hesaplanmış, yorum metnindeki her sayısal değer "
        f"hesaplama çıktısıyla doğrulanmıştır. Analiz: {content.analysis_label}."
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
